from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"D:\Code Project\Drone\training\output_20260912_multichannel\五通道压力障碍物识别模型阶段测试报告.docx")
FONT = "Microsoft YaHei"
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = RGBColor(35, 42, 52)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=130, bottom=85, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for index, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_font(run, 9.2, True, RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            run = paragraph.add_run(str(value))
            set_font(run, 9.2, False, TEXT)
    return table


def add_numbered_step(doc, number, title, body):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.45)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    title_run = paragraph.add_run(title + "  ")
    set_font(title_run, 10.5, True, TEXT)
    body_run = paragraph.add_run(body)
    set_font(body_run, 10.5, False, TEXT)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10)
normal.font.color.rgb = TEXT
normal.paragraph_format.line_spacing = 1.22
normal.paragraph_format.space_after = Pt(5)

title_style = styles["Title"]
title_style.font.name = FONT
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title_style.font.size = Pt(21)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(7)

title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (("Heading 1", 14, 8, 5), ("Heading 2", 11.5, 7, 4)):
    style = styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

doc.core_properties.title = "五通道压力障碍物识别模型阶段测试报告"
doc.core_properties.subject = "模型训练与测试结果"
doc.core_properties.author = ""

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("五通道压力障碍物识别模型阶段测试报告")
title_direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if title_direct_border is not None:
    title._p.get_or_add_pPr().remove(title_direct_border)

date = doc.add_paragraph()
date.paragraph_format.space_after = Pt(11)
date_run = date.add_run("2026年9月12日")
set_font(date_run, 10, False, RGBColor(100, 108, 118))

opening = doc.add_paragraph()
opening.paragraph_format.space_after = Pt(7)
lead = opening.add_run("阶段结论  ")
set_font(lead, 11, True, TEXT)
rest = opening.add_run("5路压力数据能够识别障碍物，但当前模型还不适合直接参与运动控制。留出测试集的实验级准确率为76.54%，障碍物检出率为72.22%。下一阶段应先完成Atlas板端联调，再通过新增独立批次改善跨批次稳定性。")
set_font(rest, 11, False, TEXT)

h = doc.add_paragraph("主要结果", style="Heading 1")
set_keep(h, True)
add_table(doc,
          ["指标", "结果", "说明"],
          [
              ["实验级准确率", "76.54%", "81次留出测试"],
              ["障碍物检出率", "72.22%", "36次有障碍实验中检出26次"],
              ["无障碍识别率", "80.00%", "45次无障碍实验中正确36次"],
              ["误报", "9次", "无障碍实验被判为有障碍"],
              ["漏检", "10次", "有障碍实验未被检出"],
          ],
          [5.2, 3.0, 8.0])

h = doc.add_paragraph("数据情况", style="Heading 1")
set_keep(h, True)
p = doc.add_paragraph("本次共使用388次有效实验。训练、验证和测试均按完整采集批次划分，同一批次不会同时出现在训练集和测试集中。")
set_keep(p)
add_table(doc,
          ["数据用途", "无障碍", "有障碍", "合计"],
          [["训练", "70", "156", "226"], ["验证", "56", "25", "81"],
           ["测试", "45", "36", "81"], ["合计", "171", "217", "388"]],
          [6.0, 3.4, 3.4, 3.4])

p = doc.add_paragraph("空文件、雷达空数据和传感器有效率不足的实验未纳入训练。")
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(0)
for run in p.runs:
    set_font(run, 9, False, RGBColor(100, 108, 118))
set_keep(p)

doc.add_page_break()

h = doc.add_paragraph("结果判断", style="Heading 1")
set_keep(h, True)
p = doc.add_paragraph("模型在验证数据上的结果明显高于测试数据，说明不同采集批次之间存在波形差异。当前结果可以证明5路压力信号具有辨识能力，但不能代表更换日期、安装位置或环境后仍能保持相同准确率。")
set_keep(p)
p = doc.add_paragraph("现阶段模型适合用于板端推理验证和误判分析。继续推进Atlas部署和数据补充；模型达到稳定应用要求前，压力识别结果作为辅助判断使用，不接管滑台运动控制。")
set_keep(p)

h = doc.add_paragraph("模型配置", style="Heading 1")
set_keep(h, True)
add_table(doc,
          ["项目", "配置"],
          [
              ["输入信号", "5路原始压力数据"],
              ["采样率", "200 Hz"],
              ["判断窗口", "最近1秒，共200点"],
              ["判断阈值", "0.87"],
              ["模型规模", "6865个可训练参数"],
              ["板端输入", "float32，形状为 1 × 5 × 1 × 200"],
          ],
          [5.2, 11.0])

p = doc.add_paragraph("数据标准化和差分处理已经包含在模型内部。Atlas端按通道1至通道5的顺序输入原始压力值，不需要重复扣除基线或标准化。")
set_keep(p)

h = doc.add_paragraph("下一阶段安排", style="Heading 1")
set_keep(h, True)
add_numbered_step(doc, 1, "完成板端联调", "将模型转换为昇腾OM格式，确认Atlas连续推理速度和输出稳定性。")
add_numbered_step(doc, 2, "复核误判数据", "逐一检查9次误报和10次漏检，确认异常来自环境变化、传感器状态还是波形重叠。")
add_numbered_step(doc, 3, "补充独立批次", "增加不同日期、风扇状态和轻微安装偏差下的数据，保证每种工况都有独立测试批次。")
add_numbered_step(doc, 4, "重新评估", "使用未参与训练的新数据统计准确率、误报率、漏检率和提前判断时间。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(4)
footer_run = footer.add_run("五通道压力障碍物识别模型阶段测试报告")
set_font(footer_run, 8.5, False, RGBColor(120, 126, 134))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

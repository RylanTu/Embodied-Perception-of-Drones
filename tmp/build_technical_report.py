from pathlib import Path
import re, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:/Code Project/Drone')
QA = ROOT / 'tmp/report_word_qa'
QA.mkdir(parents=True, exist_ok=True)
OUT = ROOT / 'docs/无人机具身气压感知实验平台技术报告.docx'
font_path = 'C:/Windows/Fonts/msyh.ttc'

def diagram(kind):
    im=Image.new('RGB',(1600,650 if kind=='system' else 460),'white')
    d=ImageDraw.Draw(im)
    font=ImageFont.truetype(font_path,29)
    small=ImageFont.truetype(font_path,25)
    def box(x,y,w,h,lines):
        d.rounded_rectangle((x,y,x+w,y+h),radius=8,fill='#F5F5F5',outline='#555555',width=2)
        text='\n'.join(lines)
        bounds=d.multiline_textbbox((0,0),text,font=font,spacing=10,align='center')
        d.multiline_text((x+(w-bounds[2])/2,y+(h-(bounds[3]-bounds[1]))/2-bounds[1]),text,font=font,fill='black',spacing=10,align='center')
    def arrow(points,label='',pos=None):
        d.line(points,fill='#555555',width=3)
        x,y=points[-1]; px,py=points[-2]; a=math.atan2(y-py,x-px)
        d.polygon([(x,y),(x-14*math.cos(a-.45),y-14*math.sin(a-.45)),(x-14*math.cos(a+.45),y-14*math.sin(a+.45))],fill='#555555')
        if label: d.text(pos,label,font=small,fill='#333333')
    if kind=='system':
        box(20,230,240,110,['SDP3x','5 路差压传感器'])
        box(350,230,250,110,['TCA9548A','通道 0～4'])
        box(690,230,240,110,['ESP32-S3','采集与运动控制'])
        box(1090,230,240,110,['Atlas','调度与数据处理'])
        box(690,30,240,100,['P100S 滑台'])
        box(1090,30,240,100,['电脑浏览器'])
        box(1090,500,240,100,['RPLIDAR'])
        box(1370,230,210,110,['CSV / SVG','识别结果'])
        arrow([(260,285),(350,285)],'I²C',(278,246))
        arrow([(600,285),(690,285)],'I²C',(617,246))
        arrow([(930,285),(1090,285)],'USB',(975,245))
        arrow([(810,230),(810,130)],'GPIO 4～7',(835,165))
        arrow([(1210,130),(1210,230)],'以太网',(1230,163))
        arrow([(1210,500),(1210,340)],'USB',(1230,414))
        arrow([(1330,285),(1370,285)])
        d.text((30,550),'压力与滑台数据经 ESP32 上传；雷达使用独立 USB 接口。',font=small,fill='#444444')
    else:
        box(30,70,260,110,['220 V AC','交流输入'])
        box(400,70,300,110,['AC/DC 电源','24 V DC'])
        box(820,70,680,110,['XL4015 级联降压网络','提供 12 V、5 V、3.3 V'])
        arrow([(290,125),(400,125)])
        arrow([(700,125),(820,125)])
        box(400,275,300,100,['独立可调电源'])
        box(820,275,300,100,['风扇'])
        arrow([(700,325),(820,325)],'18 V',(725,282))
        d.text((35,407),'低压系统共地；具体负载接点见电源接线记录。',font=small,fill='#444444')
    path=QA/(kind+'.png'); im.save(path); return path

doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297)
sec.top_margin=Mm(21); sec.bottom_margin=Mm(20)
sec.left_margin=Mm(23); sec.right_margin=Mm(23)
sec.footer_distance=Mm(10)

def set_font(obj,name='宋体',size=11):
    obj.font.name='Times New Roman'; obj.font.size=Pt(size); obj.font.color.rgb=RGBColor(0,0,0)
    rp=obj.element.get_or_add_rPr() if hasattr(obj,'element') else obj._element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rp.insert(0,rf)
    rf.set(qn('w:eastAsia'),name)

for name,size,east in [('Normal',11,'宋体'),('Title',22,'黑体'),('Subtitle',12,'宋体'),('Heading 1',15,'黑体'),('Heading 2',12,'黑体'),('Caption',9,'宋体')]:
    st=doc.styles[name]; set_font(st,east,size); st.font.italic=False
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.22
    if name.startswith('Heading'):
        st.paragraph_format.space_before=Pt(13 if name=='Heading 1' else 9)
        st.paragraph_format.keep_with_next=True
        st.font.bold=True
normal=doc.styles['Normal']; normal.paragraph_format.widow_control=True
for st in doc.styles:
    for el in list(st.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
doc.core_properties.title='无人机具身气压感知实验平台技术报告'
doc.core_properties.subject='供电系统 硬件接口 软件实现'
doc.core_properties.author=''; doc.core_properties.last_modified_by=''
doc.core_properties.comments=''
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)

def inline(p,text,size=None):
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text):
        if not part: continue
        code=part.startswith('`'); bold=part.startswith('**')
        r=p.add_run(part[1:-1] if code else part[2:-2] if bold else part)
        if size: r.font.size=Pt(size)
        if bold: r.bold=True
        if code: r.font.name='Consolas'; r.font.size=Pt(size or 10)

def add_figure(kind,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(diagram(kind)),width=Mm(163))
    p=doc.add_paragraph(caption,'Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def table(lines):
    rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',v) for v in r)]
    n=len(rows[0]); t=doc.add_table(rows=0,cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if n==2: widths=[65,99]
    elif any('文件' in x for x in rows[0]): widths=[65,99]
    else: widths=[40,57,67]
    for c,w in zip(t.columns,widths): c.width=Mm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',75),('bottom',75),('left',90),('right',90)]:
        el=OxmlElement('w:'+side); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); margins.append(el)
    pr.append(margins)
    for ix,row in enumerate(rows):
        rr=t.add_row(); trpr=rr._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ix==0: trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,txt) in enumerate(zip(rr.cells,row)):
            cell.width=Mm(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.12
            p.paragraph_format.keep_with_next=ix<len(rows)-1
            if n==2 and len(txt)<25 and j==1: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,txt,9.5)
            if ix==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E7E7'); cell._tc.get_or_add_tcPr().append(sh)
                for r in p.runs:r.bold=True
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1; p.add_run().font.size=Pt(3)

lines=(ROOT/'docs/当前供电硬件连接与软件技术报告.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
        table(group);continue
    if line.startswith('```'):
        lang=line[3:]; block=[]; i+=1
        while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
        text='\n'.join(block)
        if lang=='mermaid': add_figure('system','图 1  系统组成与数据连接')
        elif '220 V AC' in text: add_figure('power','图 2  供电系统结构')
        else:
            for k,l in enumerate(block):
                p=doc.add_paragraph(); p.paragraph_format.left_indent=Mm(5); p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_with_next=k<len(block)-1
                r=p.add_run(l); r.font.size=Pt(10); r.font.name='Consolas'
        i+=1;continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(12)
    elif line.startswith('## '):
        title=re.sub(r'^(\d+)\.\s*',r'\1  ',line[3:]).replace('：',' ')
        doc.add_paragraph(title,'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif i<6:
        p=doc.add_paragraph(line,'Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    else:
        p=doc.add_paragraph(); inline(p,line[2:] if line.startswith('- ') else line)
        if line.startswith('- '):p.paragraph_format.left_indent=Mm(4)
        else:p.paragraph_format.first_line_indent=Pt(22)
    i+=1
doc.save(OUT)
print(OUT)
